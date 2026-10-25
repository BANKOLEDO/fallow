import csv
from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

def csv_to_docx(csv_path, out_path):
    doc = Document()
    doc.add_heading('Fallow Results', 0)
    with open(csv_path, newline='', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    doc.add_paragraph(f'Rows: {len(rows)}')
    table = doc.add_table(rows=1, cols=len(rows[0].keys()) if rows else 0)
    hdr_cells = table.rows[0].cells
    for i,k in enumerate(rows[0].keys()):
        hdr_cells[i].text = k
    for r in rows:
        row_cells = table.add_row().cells
        for i,k in enumerate(rows[0].keys()):
            row_cells[i].text = str(r.get(k,''))
    doc.save(out_path)

def csv_to_pdf(csv_path, out_path):
    c = canvas.Canvas(str(out_path), pagesize=letter)
    w,h = letter
    y = h-40
    c.setFont('Helvetica-Bold', 10)
    with open(csv_path, newline='', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        keys = list(rows[0].keys()) if rows else []
        x = 40
        for k in keys:
            c.drawString(x,y,k[:12])
            x += 90
        y -= 15
        c.setFont('Helvetica', 7)
        for r in rows[:200]:  # limit for readability
            x = 40
            for k in keys:
                c.drawString(x,y,str(r.get(k,''))[:14])
                x += 90
            y -= 10
            if y < 40:
                c.showPage()
                y = h-40
                c.setFont('Helvetica', 7)
    c.save()

if __name__ == '__main__':
    csv_to_docx('fallow_predictions.csv', 'fallow_predictions.docx')
    csv_to_pdf('fallow_predictions.csv', 'fallow_predictions.pdf')
    csv_to_docx('fallow_data.csv', 'fallow_data.docx')
    csv_to_pdf('fallow_data.csv', 'fallow_data.pdf')
    print('done')
